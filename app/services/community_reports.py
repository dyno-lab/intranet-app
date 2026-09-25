"""Read-only report data: attendance determines participation; fiscal copies supply demographics."""
from __future__ import annotations

from collections import defaultdict
from datetime import date
import json
from io import BytesIO
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import and_, select
from sqlalchemy.orm import Session

from app.helpers.reports import AGE_BUCKETS, get_age_bucket
from app.models.community import CPFiscalYear, CPProgram
from app.models.community_activity import CPActivity, CPADMServiceType, CPADMServiceActivity
from app.models.community_fiscal import CPFiscalParticipant
from app.models.community_operations import CPActivitySession, CPAttendance, CPGradeReport, CPGradeItem

REPORT_TYPES = {
    "bonafide": "Bonafide", "no-duplicados": "Participantes no duplicados",
    "participaciones": "Participaciones", "por-programa": "Desglose por programa",
    "adm": "Informe ADM", "notas": "Notas escolares",
}


def age_at(born, reference: date):
    if not born:
        return None
    born = date.fromisoformat(born) if isinstance(born, str) else born
    return reference.year - born.year - ((reference.month, reference.day) < (born.month, born.day))


def _name(snapshot):
    return " ".join(str(snapshot.get(key) or "") for key in ("nombre", "inicial", "apellido_paterno", "apellido_materno")).strip()


def _demographics(snapshots, reference):
    buckets = {key: [label, 0, 0, 0] for key, label in AGE_BUCKETS}
    buckets["unknown"] = ["Sin fecha de nacimiento", 0, 0, 0]
    for snapshot in snapshots:
        bucket = get_age_bucket(age_at(snapshot.get("fecha_nacimiento"), reference)) or "unknown"
        row = buckets[bucket]
        gender = str(snapshot.get("genero") or "").upper()
        row[1] += int(gender == "F")
        row[2] += int(gender == "M")
        row[3] += 1
    return list(buckets.values())


def _adm_demographics(snapshots, reference):
    limits = ((0, 5), (6, 11), (12, 17), (18, 21), (22, 25), (26, 45), (46, 59), (60, 74), (75, 9999))
    rows = [[f"{low}–{high}" if high < 9999 else "75+", 0, 0, 0, 0, 0] for low, high in limits]
    rows.append(["Sin edad válida", 0, 0, 0, 0, 0])
    people = list(snapshots)
    for snapshot in people:
        age = age_at(snapshot.get("fecha_nacimiento"), reference)
        index = next((i for i, (low, high) in enumerate(limits) if age is not None and low <= age <= high), len(limits))
        row = rows[index]
        row[1] += int(snapshot.get("genero") == "F")
        row[2] += int(snapshot.get("genero") == "M")
        row[3] += 1
        row[5] += int(str(snapshot.get("vca") or "").upper() in {"SI", "SÍ"})
    for row in rows:
        row[4] = round(row[3] * 100 / len(people), 2) if people else 0
    return rows


def build_report(db: Session, *, report_type: str, fiscal_year_id: int,
                 program_ids: set[int], start_date: date, end_date: date) -> dict:
    if report_type not in REPORT_TYPES:
        raise ValueError("Seleccione un reporte válido.")
    fiscal = db.get(CPFiscalYear, fiscal_year_id)
    if fiscal is None:
        raise ValueError("Año fiscal no encontrado.")
    if start_date > end_date or start_date < fiscal.start_date or end_date > fiscal.end_date:
        raise ValueError("El período debe estar dentro del año fiscal seleccionado.")
    if not program_ids:
        raise ValueError("Seleccione al menos un programa.")
    programs = {p.program_id: p for p in db.scalars(select(CPProgram).where(CPProgram.program_id.in_(program_ids)))}
    if set(programs) != program_ids:
        raise ValueError("Programa no encontrado.")
    session_filter = (
        CPActivitySession.fiscal_year_id == fiscal_year_id,
        CPActivitySession.program_id.in_(program_ids),
        CPActivitySession.session_date >= start_date,
        CPActivitySession.session_date <= end_date,
    )
    # Outer join makes missing historical copies an explicit error, not silent data loss.
    records = db.execute(select(CPActivitySession, CPAttendance, CPFiscalParticipant, CPActivity).join(
        CPAttendance, CPAttendance.session_id == CPActivitySession.session_id
    ).outerjoin(CPFiscalParticipant, and_(CPFiscalParticipant.participant_id == CPAttendance.participant_id,
                                        CPFiscalParticipant.fiscal_year_id == CPActivitySession.fiscal_year_id)
    ).join(CPActivity, CPActivity.activity_id == CPActivitySession.activity_id).where(
        *session_filter, CPAttendance.is_present == True  # noqa: E712
    ).order_by(CPActivitySession.session_date, CPActivitySession.session_id, CPAttendance.participant_id)).all()
    snapshots, by_program, participant_programs = {}, defaultdict(set), defaultdict(set)
    for session, attendance, historical, activity in records:
        if historical is None:
            raise ValueError("Hay participaciones sin datos del año fiscal. Revise la sincronización antes de generar el informe.")
        snapshots[attendance.participant_id] = json.loads(historical.snapshot_json)
        by_program[session.program_id].add(attendance.participant_id)
        participant_programs[attendance.participant_id].add(programs[session.program_id].code)
    result = {
        "title": REPORT_TYPES[report_type], "type": report_type, "fiscal": fiscal,
        "program_label": ", ".join(p.code for p in programs.values()),
        "start_date": start_date, "end_date": end_date,
        "unique_count": len(snapshots), "attendance_count": len(records),
        "headers": [], "rows": [], "sections": [],
    }
    if report_type == "bonafide":
        result["headers"] = ["Expediente", "Participante", "Género", "Edad", "Pueblo", "Programas con participación", "Dirección física", "Primera vez"]
        result["rows"] = [[s.get("expediente_num", ""), _name(s), s.get("genero", ""),
                           age_at(s.get("fecha_nacimiento"), end_date), s.get("pueblo", ""),
                           ", ".join(sorted(participant_programs[pid])), s.get("direccion_fisica", ""), s.get("primera_vez", "")]
                          for pid, s in sorted(snapshots.items(), key=lambda pair: _name(pair[1]))]
    elif report_type in {"no-duplicados", "participaciones"}:
        result["headers"] = ["Grupo de edad", "Femenino", "Masculino", "Total"]
        people = snapshots.values() if report_type == "no-duplicados" else [snapshots[a.participant_id] for _, a, _, _ in records]
        result["rows"] = _demographics(people, end_date)
        if report_type == "participaciones":
            result["sections"].append({"title": "Detalle de participaciones", "headers": ["Fecha", "Programa", "Actividad", "Expediente", "Participante"],
                                       "rows": [[s.session_date.isoformat(), programs[s.program_id].code, activity.code,
                                                 snapshots[a.participant_id].get("expediente_num", ""), _name(snapshots[a.participant_id])]
                                                for s, a, _, activity in records]})
    elif report_type == "por-programa":
        result["headers"] = ["Programa", "Participantes no duplicados", "Participaciones"]
        result["rows"] = [[program.code, len(by_program[pid]), sum(s.program_id == pid for s, _, _, _ in records)] for pid, program in programs.items()]
        result["rows"].append(["Total consolidado sin duplicados", len(snapshots), len(records)])
    elif report_type == "adm":
        services = db.scalars(select(CPADMServiceType).where(
            CPADMServiceType.fiscal_year_id == fiscal_year_id, CPADMServiceType.program_id.in_(program_ids)
        ).order_by(CPADMServiceType.program_id, CPADMServiceType.sort_order, CPADMServiceType.name)).all()
        mappings = db.scalars(select(CPADMServiceActivity).where(
            CPADMServiceActivity.fiscal_year_id == fiscal_year_id,
            CPADMServiceActivity.program_id.in_(program_ids), CPADMServiceActivity.is_active == True  # noqa: E712
        )).all()
        mapping = {(m.program_id, m.activity_id): m.adm_service_type_id for m in mappings}
        all_sessions = db.scalars(select(CPActivitySession).where(*session_filter)).all()
        adm_records = [(s, a, h, activity) for s, a, h, activity in records if (s.program_id, s.activity_id) in mapping]
        adm_snapshots = {a.participant_id: snapshots[a.participant_id] for _, a, _, _ in adm_records}
        result["unique_count"], result["attendance_count"] = len(adm_snapshots), len(adm_records)
        result["headers"] = ["Programa", "Tipo de servicio", "Servicios realizados", "Participaciones", "Participantes no duplicados"]
        for service in services:
            service_sessions = {s.session_id for s in all_sessions if mapping.get((s.program_id, s.activity_id)) == service.adm_service_type_id}
            attended = [(s, a) for s, a, _, _ in records if s.session_id in service_sessions]
            result["rows"].append([programs[service.program_id].code, service.name, len(service_sessions), len(attended), len({a.participant_id for _, a in attended})])
        result["sections"].append({"title": "Participantes por edad y género", "headers": ["Edad", "Femenino", "Masculino", "Total", "%", "VCA"], "rows": _adm_demographics(adm_snapshots.values(), end_date)})
        for key, title in (("composicion_familiar", "Composición familiar"), ("fuente_ingreso_principal", "Fuente de ingreso"), ("rango_ingreso", "Rango de ingreso")):
            counts = defaultdict(int)
            for snapshot in adm_snapshots.values():
                counts[snapshot.get(key) or "Sin registrar"] += 1
            result["sections"].append({"title": title, "headers": ["Categoría", "Participantes"], "rows": [[key, value] for key, value in sorted(counts.items())]})
        unclassified = [s for s in all_sessions if (s.program_id, s.activity_id) not in mapping]
        if unclassified:
            result["sections"].append({"title": "Sesiones sin clasificación ADM (fuera del total ADM)", "headers": ["Fecha", "Programa", "Actividad", "Participaciones"], "rows": [
                [s.session_date.isoformat(), programs[s.program_id].code, db.get(CPActivity, s.activity_id).code,
                 sum(record.session_id == s.session_id for record, _, _, _ in records)] for s in unclassified]})
    elif report_type == "notas":
        from app.services.community_operations import grade_letter
        result["headers"] = ["Programa", "Período", "Expediente", "Participante", "Grado", "Español", "Inglés", "Matemáticas", "Ciencias", "Estudios sociales", "Electiva 1", "Electiva 2", "Electiva 3", "Electiva 4", "Promedio", "Calificación", "Salón contenido"]
        grade_participants = set()
        grades = db.execute(select(CPGradeReport, CPGradeItem).join(CPGradeItem).where(
            CPGradeReport.fiscal_year_id == fiscal_year_id, CPGradeReport.program_id.in_(program_ids)
        ).order_by(CPGradeReport.report_year, CPGradeReport.report_month)).all()
        for report, item in grades:
            # Monthly academic results are selected by their reporting month.
            if not (start_date.year, start_date.month) <= (report.report_year, report.report_month) <= (end_date.year, end_date.month):
                continue
            if item.participant_id not in by_program[report.program_id]:
                continue
            snapshot = snapshots[item.participant_id]
            grade_participants.add((report.program_id, item.participant_id))
            result["rows"].append([programs[report.program_id].code, f"{report.report_year}-{report.report_month:02d}",
                                   snapshot.get("expediente_num", ""), _name(snapshot), item.grade_level or "",
                                   *[getattr(item, key) for key in ("spanish_grade", "english_grade", "math_grade", "science_grade", "social_studies_grade", "elective_1_grade", "elective_2_grade", "elective_3_grade", "elective_4_grade", "average_grade")], grade_letter(item.average_grade), "Sí" if item.is_content_room else "No"])
        result["unique_count"] = len({pid for _, pid in grade_participants})
        result["attendance_count"] = sum((s.program_id, a.participant_id) in grade_participants for s, a, _, _ in records)
    return result


def excel_bytes(report: dict) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    sections = [{"title": report["title"], "headers": report["headers"], "rows": report["rows"]}, *report["sections"]]
    for index, section in enumerate(sections, 1):
        sheet = workbook.create_sheet(f"{index}. {section['title']}"[:31])
        sheet.append(["Comunidad y Prevención", report["fiscal"].code])
        sheet.append([report["program_label"], f"{report['start_date']} — {report['end_date']}"])
        sheet.append(["Participantes no duplicados", report["unique_count"], "Participaciones", report["attendance_count"]])
        sheet.append(section["headers"])
        for row in section["rows"]:
            sheet.append(["" if value is None else value for value in row])
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"  # Keep user-entered formula-like text literal.
        for cell in sheet[4]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="126443")
        sheet.freeze_panes = "A5"
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(42, max(14, max(len(str(c.value or "")) for c in column) + 2))
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def pdf_bytes(report: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, CondPageBreak
    buffer = BytesIO()
    width, height = landscape(letter)
    doc = SimpleDocTemplate(buffer, pagesize=(width, height), leftMargin=28, rightMargin=28, topMargin=30, bottomMargin=35)
    styles = getSampleStyleSheet()
    small = ParagraphStyle("CP cell", parent=styles["Normal"], fontSize=7, leading=9)
    heading = ParagraphStyle("CP section", parent=styles["Heading3"], keepWithNext=False)
    header = ParagraphStyle("CP column", parent=small, fontName="Helvetica-Bold")
    story = [Paragraph("Comunidad y Prevención", styles["Title"]), Paragraph(escape(report["title"]), styles["Heading2"]),
             Paragraph(escape(f"{report['fiscal'].code} · {report['program_label']} · {report['start_date']} a {report['end_date']}"), styles["Normal"]),
             Paragraph(f"Participantes no duplicados: {report['unique_count']} · Participaciones: {report['attendance_count']}", styles["Normal"]), Spacer(1, 12)]
    for section in [{"title": "Resultados", "headers": report["headers"], "rows": report["rows"]}, *report["sections"]]:
        section_heading = Paragraph(escape(section["title"]), heading)
        rows = section["rows"] or [["Sin datos para esta selección"] + [""] * (len(section["headers"]) - 1)]
        headers = section["headers"]
        weights = [1] * len(headers)
        if report["type"] == "notas":
            headers = ["Programa", "Período", "Expediente", "Participante", "Grado", "Esp.", "Ing.", "Mat.", "Cien.", "Est. soc.", "Elec. 1", "Elec. 2", "Elec. 3", "Elec. 4", "Prom.", "Letra", "S. cont."]
            weights = [8, 9, 13, 24, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6, 7, 6, 7]
        elif report["type"] == "bonafide":
            weights = [12, 23, 7, 6, 10, 15, 22, 8]
        cells = [[Paragraph(escape(str(value)), header) for value in headers]]
        cells.extend([[Paragraph(escape(str(value if value is not None else "")).replace("–", "-"), small) for value in row] for row in rows])
        column_widths = [(width - 56) * weight / sum(weights) for weight in weights]
        table = Table(cells, colWidths=column_widths, repeatRows=1, hAlign="LEFT")
        table_style = TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#deeee4")), ("GRID", (0, 0), (-1, -1), .3, colors.HexColor("#b7cbbf")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)])
        table.setStyle(table_style)
        first_rows = Table(cells[:2], colWidths=column_widths)
        first_rows.setStyle(table_style)
        minimum_height = first_rows.wrap(width - 56, height)[1] + section_heading.wrap(width - 56, height)[1] + 24
        # Keep a section title with its header and first result, while letting a
        # long table start on the current page instead of moving it in full.
        story.extend([CondPageBreak(minimum_height), section_heading, table, Spacer(1, 12)])
    def footer(canvas, document):
        canvas.setFont("Helvetica", 8)
        canvas.drawRightString(width - 28, 18, f"Página {document.page}")
        if document.page > 1:
            canvas.drawString(28, height - 18, f"Comunidad y Prevención · {report['title']} · {report['fiscal'].code}")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()
