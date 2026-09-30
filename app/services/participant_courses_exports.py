"""Presentation only: the same saved monthly choices in PDF, print and Excel."""
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import LongTable, PageBreak, Paragraph, SimpleDocTemplate, Spacer, TableStyle

from app.services.full_monthly_report_charts import _decorate

ASSETS = Path(__file__).resolve().parents[1] / 'static' / 'reports' / 'full_monthly'
TITLE = 'Informe de cursos - Actividad 2.b.5'


def course_report_pdf(data):
    output = BytesIO()
    width, height = landscape(letter)
    doc = SimpleDocTemplate(output, pagesize=(width, height), leftMargin=48, rightMargin=48,
                            topMargin=182, bottomMargin=155, title=TITLE,
                            author='Centros Sor Isolina Ferré - Faro de Esperanza')
    normal = ParagraphStyle('CourseCell', fontName='Helvetica', fontSize=9, leading=11)
    bold = ParagraphStyle('CourseHeading', parent=normal, fontName='Helvetica-Bold')
    def paragraph(value, heading=False):
        return Paragraph(escape(str(value or '')), bold if heading else normal)
    def page(canvas, document):
        canvas.saveState()
        # Exact source assets, dimensions and placement from the approved full report.
        _decorate(canvas, (width, height), reference='pregnancy')
        canvas.setFont('Helvetica-Bold', 13)
        canvas.drawCentredString(width / 2, 493, TITLE)
        meta = paragraph(f"Período: {data['period_label']} | Residencial: {data['residential_name']}")
        _, h = meta.wrap(width - 96, 48)
        meta.drawOn(canvas, 48, 479 - h)
        canvas.setFont('Helvetica', 8)
        canvas.drawRightString(width - 48, 48, f'Página {document.page}')
        canvas.restoreState()
    story = []
    # Keep month labels legible on letter paper. Wider periods continue in groups
    # of three months; people are never recounted in the report's unique total.
    for start in range(0, len(data['months']), 3):
        if start:
            story.append(PageBreak())
        story.extend([paragraph(f"Propuestas: {data['proposal_label']}"), Spacer(1, 8)])
        months = data['months'][start:start + 3]
        headers = ['Participante', 'Residencial', *[m['label'] for m in months]]
        rows = [[paragraph(v, True) for v in headers]]
        for row in data['rows']:
            rows.append([paragraph(row['name']), paragraph(row['residential_name']),
                         *[paragraph(c['label']) for c in row['cells'][start:start + 3]]])
        if not data['rows']:
            rows.append([paragraph('Sin participantes con asistencia en el período.'), *[''] * (len(headers) - 1)])
        rows.append([paragraph(f"Participantes únicos en todo el período: {data['total']}", True),
                     *[''] * (len(headers) - 1)])
        table = LongTable(rows, colWidths=[175, 130, *([391 / len(months)] * len(months))], repeatRows=1)
        table.setStyle(TableStyle([
            ('GRID', (0, 0), (-1, -1), .5, colors.HexColor('#555555')),
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#f1f1f1')),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('LEFTPADDING', (0, 0), (-1, -1), 6), ('RIGHTPADDING', (0, 0), (-1, -1), 6),
            ('TOPPADDING', (0, 0), (-1, -1), 6), ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('SPAN', (0, -1), (-1, -1)), ('NOSPLIT', (0, -2), (-1, -1)),
        ]))
        story.append(table)
    doc.build(story, onFirstPage=page, onLaterPages=page)
    return output.getvalue()


def course_report_excel(data):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Cursos 2.b.5'
    columns = 3 + len(data['months'])
    last = get_column_letter(columns)
    logo = Image(str(ASSETS / 'chart_133_image16.png'))
    logo.width, logo.height = 252, 85
    sheet.add_image(logo, 'A1')
    sheet.merge_cells(f'A5:{last}5')
    sheet['A5'] = TITLE
    sheet['A5'].font = Font(size=14, bold=True)
    sheet.merge_cells(f'A6:{last}6')
    sheet['A6'] = f"Propuestas: {data['proposal_label']} | Período: {data['period_label']} | Residencial: {data['residential_name']}"
    sheet['A6'].data_type = 's'
    sheet['A6'].alignment = Alignment(wrap_text=True)
    sheet.row_dimensions[6].height = 40
    headers = ['Nombre del participante', 'Expediente', 'Residencial', *[m['label'] for m in data['months']]]
    for index, value in enumerate(headers, 1):
        cell = sheet.cell(8, index, value)
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='EDF2FA')
    for row_index, row in enumerate(data['rows'], 9):
        for col, value in enumerate([row['name'], row['expediente_num'], row['residential_name'], *[c['label'] for c in row['cells']]], 1):
            cell = sheet.cell(row_index, col, value)
            cell.data_type = 's'  # Names and expediente numbers are text, never formulas.
        sheet.row_dimensions[row_index].height = 42
    end = 8 + len(data['rows'])
    for row in sheet.iter_rows(min_row=8, max_row=max(end, 8), max_col=columns):
        for cell in row:
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            cell.border = Border(bottom=Side(style='thin', color='DCE3ED'))
    sheet.cell(end + 2, 1, 'Participantes únicos en todo el período')
    sheet.cell(end + 2, 2, data['total'])
    footer = Image(str(ASSETS / 'chart_133_image14.png'))
    footer.width, footer.height = 235, 94
    sheet.add_image(footer, f'A{end + 4}')
    for index in range(1, columns + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 34 if index == 1 else 22 if index == 2 else 32
    sheet.freeze_panes = 'D9'
    sheet.auto_filter.ref = f'A8:{last}{max(8, end)}'
    sheet.print_title_rows = '5:8'
    sheet.print_options.horizontalCentered = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_LETTER
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.oddFooter.right.text = 'Página &P'
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
