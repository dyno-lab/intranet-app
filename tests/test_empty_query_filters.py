import unittest
from types import SimpleNamespace
from io import BytesIO
from unittest.mock import patch
from urllib.parse import parse_qs, urlsplit

from fastapi import Depends, FastAPI, HTTPException, Request
from fastapi.routing import APIRoute
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool
from starlette.middleware.sessions import SessionMiddleware
from openpyxl import load_workbook

from tests import test_full_monthly_report_data as fixtures
from app.api.routes import admin, catalogs, consolidado_mensual_global, hoja_cotejo_admin, plantilla_duplicado, reports, ui
from app.core.auth import get_current_user, require_admin, require_admin_or_supervisor
from app.core.residential_scope import require_faro_access
from app.helpers.report_proposals import report_proposal_selection


MODULES = ((admin, "/ui/admin"), (catalogs, "/ui/admin/catalogs"),
           (consolidado_mensual_global, "/ui/admin"), (hoja_cotejo_admin, "/ui/admin"),
           (plantilla_duplicado, "/ui/admin"), (reports, "/ui/reports"), (ui, "/ui"))


class EmptyQueryContractTests(unittest.TestCase):
    def test_all_optional_integer_query_filters_accept_blank_but_reject_invalid(self):
        checked = 0
        for module, prefix in MODULES:
            for route in module.router.routes:
                if not isinstance(route, APIRoute) or "GET" not in route.methods:
                    continue
                for field in route.dependant.query_params:
                    if field.default is not None or "int" not in str(field.field_info.annotation):
                        continue
                    checked += 1
                    with self.subTest(route=prefix + route.path, field=field.name):
                        for raw in ("", " ", "\t"):
                            value, errors = field.validate(raw, {}, loc=("query", field.name))
                            self.assertFalse(errors)
                            self.assertIsNone(value)
                        for raw, expected in (("0", 0), ("-2", -2), ("12", 12)):
                            value, errors = field.validate(raw, {}, loc=("query", field.name))
                            self.assertFalse(errors)
                            self.assertEqual(value, expected)
                        for raw in ("abc", "1.5"):
                            self.assertTrue(field.validate(raw, {}, loc=("query", field.name))[1])
        self.assertGreaterEqual(checked, 125)

    def test_repeated_proposals_ignore_only_empty_choices(self):
        for query, expected in (("proposal_id=", None),
                                ("proposal_id=1&proposal_id=&proposal_id=2", [1, 2]),
                                ("proposal_id=%20&proposal_id=2&proposal_id=2", 2)):
            request = Request({"type": "http", "query_string": query.encode()})
            self.assertEqual(report_proposal_selection(request), expected)
        for invalid in ("abc", "0", "-1", "1.5"):
            request = Request({"type": "http", "query_string": f"proposal_id=1&proposal_id={invalid}".encode()})
            with self.assertRaises(HTTPException):
                report_proposal_selection(request)


class EmptyQueryScreenTests(unittest.TestCase):
    def setUp(self):
        self.source = fixtures.FullMonthlyReportDataTests()
        with patch.object(fixtures, "create_engine", side_effect=lambda *_args: create_engine(
                "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)):
            self.source.setUp()
        self.addCleanup(self.source.doCleanups)
        self.user = SimpleNamespace(user_id=90, username="Prueba", role="admin", residential_id=None)
        self.app = FastAPI()
        self.app.add_middleware(SessionMiddleware, secret_key="empty-query-filter-test")
        for module, prefix in MODULES:
            self.app.include_router(module.router, prefix=prefix, dependencies=[Depends(require_faro_access)])
        def db():
            with Session(self.source.engine) as session:
                yield session
        self.app.dependency_overrides[reports.get_db] = db
        for dependency in (get_current_user, require_admin, require_admin_or_supervisor):
            self.app.dependency_overrides[dependency] = lambda: self.user
        permission = patch("app.core.residential_scope._FARO_PERMISSION_DEPENDENCY", return_value=self.user)
        permission.start()
        self.addCleanup(permission.stop)
        # The bundled Starlette dropped the positional TemplateResponse API
        # used by these existing screens. Adapt that API in the harness while
        # rendering the real templates; this fix does not change their layout.
        for module, _ in MODULES:
            original_response = module.templates.TemplateResponse
            def render(*args, _original=original_response, **kwargs):
                if args and isinstance(args[0], str):
                    name, context = args[:2]
                    return _original(request=context["request"], name=name, context=context, **kwargs)
                return _original(*args, **kwargs)
            renderer = patch.object(module.templates, "TemplateResponse", side_effect=render)
            renderer.start()
            self.addCleanup(renderer.stop)
        self.client = TestClient(self.app)
        self.addCleanup(self.client.close)

    def test_administration_blank_proposal_shows_the_normal_screen(self):
        for path in ("activity-codes", "report-templates", "visits", "adm", "vca", "proposal-participants", "report-programs"):
            with self.subTest(path=path):
                response = self.client.get(f"/ui/admin/{path}?proposal_id=")
                self.assertEqual(response.status_code, 200, response.text)

    def test_report_previews_accept_empty_filters(self):
        for path in ("", "bonafide", "no-duplicado", "duplicado", "visitas", "notas", "embarazo",
                     "desercion-escolar", "vca", "adm", "por-programa", "hoja-cotejo", "productividad"):
            with self.subTest(path=path):
                response = self.client.get(f"/ui/reports/{path}?proposal_id=&month=&year=&employee_id=")
                self.assertEqual(response.status_code, 200, response.text)

    def test_other_existing_filter_screens_keep_accepting_empty_values(self):
        for path in (
            "/ui/new-list?age_min=&age_max=&residential_id=&age_range=",
            "/ui/listado?proposal_id=&month=&year=&residential_id=&from_date=&to_date=",
            "/ui/new-list/1/expediente?participation_from_date=&participation_to_date=&participation_proposal_id=",
            "/ui/admin/catalogs?selected_type_id=",
        ):
            with self.subTest(path=path):
                self.assertEqual(self.client.get(path).status_code, 200)

    def test_custom_period_exports_preserve_multiple_proposals_and_global_scope(self):
        query = "?proposal_id=&proposal_id=1&proposal_id=2&proposal_id=&month=&year=&employee_id=0&period_type=custom&start_date=2026-07-01&end_date=2026-07-31"
        response = self.client.get("/ui/reports/bonafide/excel" + query)
        self.assertEqual(response.status_code, 200, response.text[:100])
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual([workbook["Bonafide"].cell(row=i, column=3).value for i in (9, 10)],
                         ["Participante 1 Prueba", "Participante 2 Prueba"])
        with patch.object(reports, "render_template_to_chromium_pdf_bytes", return_value=b"%PDF-test") as render:
            response = self.client.get("/ui/reports/bonafide/pdf/download" + query)
            self.assertEqual(response.status_code, 200, response.text)
            context = render.call_args.kwargs["context"]
            self.assertEqual(context["selected_proposal_ids"], [1, 2])
            self.assertEqual(context["selected_employee_id"], 0)
            self.assertEqual(len(context["rows"]), 2)

    def test_required_identifiers_are_not_treated_as_optional_filters(self):
        self.assertEqual(self.client.get("/ui/admin/report-templates/versions/download-word?report_template_id=").status_code, 422)
        self.assertEqual(self.client.get("/ui/listado/not-an-id").status_code, 422)
        self.assertEqual(self.client.post("/ui/reports/visitas/referrals/save", data={"proposal_id": "", "month": "", "year": ""}).status_code, 422)

    def test_report_center_redirects_empty_scope_without_literal_none(self):
        for report in ("bonafide", "no-duplicado", "duplicados", "visitas", "notas", "embarazo",
                       "desercion-escolar", "vca", "adm", "por-programa", "hoja-cotejo", "productividad", "todos"):
            with self.subTest(report=report):
                response = self.client.get("/ui/reports/run", params={"report_key": report,
                    "proposal_id": "", "month": "", "year": "", "employee_id": ""}, follow_redirects=False)
                self.assertEqual(response.status_code, 303, response.text)
                self.assertNotIn("None", response.headers["location"])
                self.assertEqual(self.client.get(response.headers["location"]).status_code, 200)
        response = self.client.get("/ui/reports/run?report_key=productividad&proposal_id=1&month=7&year=2026&employee_id=0", follow_redirects=False)
        self.assertEqual(parse_qs(urlsplit(response.headers["location"]).query)["employee_id"], ["0"])

    def test_custom_admin_periods_accept_empty_unused_month_and_year(self):
        for module, path in ((consolidado_mensual_global, "consolidado-mensual-global"),
                             (plantilla_duplicado, "plantilla-duplicado"), (hoja_cotejo_admin, "hoja-cotejo")):
            # Isolate the SQL Server-specific aggregate; exercise routing and period validation.
            def context(*args, **kwargs):
                module._validate_period(kwargs["period_type"], args[2], args[3], kwargs["start_date"], kwargs["end_date"])
                self.assertIsNone(args[4])
                return {}
            with self.subTest(path=path), patch.object(module, "_build_context", side_effect=context), \
                 patch.object(module.templates, "TemplateResponse", return_value=reports.HTMLResponse("OK")):
                response = self.client.get(f"/ui/admin/{path}?period_type=custom&month=&year=&proposal_id=&residential_id=&start_date=2026-07-01&end_date=2026-07-31")
                self.assertEqual(response.status_code, 200, response.text)


if __name__ == "__main__":
    unittest.main()
