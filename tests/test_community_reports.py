from __future__ import annotations

import os
import tempfile
import unittest
from datetime import date, datetime, timezone
from io import BytesIO
from pathlib import Path

os.environ.setdefault("DB_SERVER", "test-server")
os.environ.setdefault("DB_NAME", "test-db")
os.environ.setdefault("DB_USER", "test-user")
os.environ.setdefault("DB_PASSWORD", "test-password")
os.environ.setdefault("SESSION_SECRET", "community-test-session-secret-not-production")

from fastapi import FastAPI
from fastapi.staticfiles import StaticFiles
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import create_engine, event, select
from sqlalchemy.orm import Session
from starlette.middleware.sessions import SessionMiddleware

from app.api.deps import get_db
from app.api.routes import community_reports
from app.core.community_access import CommunityContext, require_community_context
from app.models.base import Base
from app.models.community import CPFiscalYear, CPParticipant, CPParticipantProgram, CPProgram, CPSequence
from app.models.community_activity import COMMUNITY_ACTIVITY_MODELS
from app.models.community_catalog import CPProfileField, CPProfileValue
from app.models.community_fiscal import CPFiscalEnrollment, CPFiscalParticipant, CPFiscalState, CPEnrollmentPeriod
from app.models.community_operations import COMMUNITY_OPERATION_MODELS, CPAttendance
from app.models.residential import Residential
from app.models.user import User
from app.services.community import create_fiscal_year, create_participant, create_program
from app.services.community_activity import create_activity, create_adm_service_type, assign_adm_activity
from app.services.community_fiscal import discharge_participant, enroll_participant, set_snapshot_freeze, sync_participants
from app.services.community_operations import create_activity_session, set_session_attendance, create_grade_report, save_grade_item
from app.services.community_reports import REPORT_TYPES, build_report, excel_bytes, pdf_bytes


class _ReportFixture(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.engine = create_engine(f"sqlite:///{Path(self.temp.name) / 'reports.db'}")

        @event.listens_for(self.engine, "connect")
        def configure(connection, _):
            connection.execute("PRAGMA foreign_keys=ON")
            connection.create_function("sysutcdatetime", 0, lambda: datetime.now(timezone.utc).isoformat())

        models = (Residential, User, CPProgram, CPFiscalYear, CPSequence, CPParticipant, CPParticipantProgram,
                  CPProfileField, CPProfileValue, CPFiscalState, CPFiscalParticipant, CPFiscalEnrollment,
                  CPEnrollmentPeriod, *COMMUNITY_ACTIVITY_MODELS, *COMMUNITY_OPERATION_MODELS)
        Base.metadata.create_all(self.engine, tables=[model.__table__ for model in models])
        self.db = Session(self.engine, expire_on_commit=False)
        self.actor = User(username="reports-test", password_hash="not-used", role="viewer", is_active=True)
        self.db.add(self.actor)
        self.db.flush()
        self.voca = create_program(self.db, "VOCA", "VOCA")
        self.tanf = create_program(self.db, "TANF-M", "TANF-M")
        self.year = create_fiscal_year(self.db, "2024", "Fiscal 2024", date(2024, 1, 15), date(2024, 12, 20))
        self.person = self.participant("María")
        self.nonattender = self.participant("Solo expediente")
        sync_participants(self.db, self.year.fiscal_year_id, [self.person.participant_id, self.nonattender.participant_id],
                          actor_user_id=self.actor.user_id)
        self.activities = {}
        for program in (self.voca, self.tanf):
            for person in (self.person, self.nonattender):
                enroll_participant(self.db, participant_id=person.participant_id, program_id=program.program_id,
                                   fiscal_year_id=self.year.fiscal_year_id, start_date=date(2024, 1, 15),
                                   actor_user_id=self.actor.user_id)
            activity = create_activity(self.db, program_id=program.program_id, code="TALLER", description="Taller",
                                        fiscal_year_ids=[self.year.fiscal_year_id])
            self.activities[program.program_id] = activity
            service = create_adm_service_type(self.db, program_id=program.program_id,
                                              fiscal_year_id=self.year.fiscal_year_id, name="Orientación")
            assign_adm_activity(self.db, service_type_id=service.adm_service_type_id, program_id=program.program_id,
                                fiscal_year_id=self.year.fiscal_year_id, activity_id=activity.activity_id)
        self.session = self.attend(self.voca, date(2024, 2, 1))
        self.attend(self.voca, date(2024, 2, 2))
        self.attend(self.tanf, date(2024, 2, 3))
        grade = create_grade_report(self.db, fiscal_year_id=self.year.fiscal_year_id, program_id=self.voca.program_id,
                                    report_year=2024, report_month=2, actor_user_id=self.actor.user_id)
        save_grade_item(self.db, report_id=grade.report_id, participant_id=self.person.participant_id,
                        fields={"grade_level": "8", "spanish_grade": 90, "english_grade": 80})
        self.db.commit()

    def participant(self, name):
        return create_participant(self.db, actor_user_id=self.actor.user_id, exp_year=2024,
            program_ids=[self.voca.program_id, self.tanf.program_id],
            fields={"nombre": name, "apellido_paterno": "Rivera", "genero": "F",
                    "fecha_nacimiento": "2010-01-01", "pueblo": "Ponce"})

    def attend(self, program, on_date):
        row = create_activity_session(self.db, fiscal_year_id=self.year.fiscal_year_id,
            program_id=program.program_id, activity_id=self.activities[program.program_id].activity_id,
            session_date=on_date, actor_user_id=self.actor.user_id)
        set_session_attendance(self.db, session_id=row.session_id, present_participant_ids=[self.person.participant_id])
        return row

    def report(self, report_type="no-duplicados", **kwargs):
        params = dict(report_type=report_type, fiscal_year_id=self.year.fiscal_year_id,
                      program_ids={self.voca.program_id, self.tanf.program_id},
                      start_date=self.year.start_date, end_date=self.year.end_date)
        params.update(kwargs)
        return build_report(self.db, **params)

    def tearDown(self):
        self.db.close()
        self.engine.dispose()
        self.temp.cleanup()


class CommunityReportDataTests(_ReportFixture):
    def test_long_grade_pdf_starts_results_on_first_page(self):
        report = self.report("notas")
        row = list(report["rows"][0])
        row[3] = "María Alejandra Rivera Fernández"
        report["rows"] = [list(row) for _ in range(30)]
        reader = PdfReader(BytesIO(pdf_bytes(report)))
        self.assertIn("María", reader.pages[0].extract_text())
        for page in reader.pages[1:]:
            self.assertIn("Comunidad y Prevención", page.extract_text())

    def test_attendance_only_and_two_program_union_is_one_person(self):
        report = self.report()
        self.assertEqual(report["unique_count"], 1)
        self.assertEqual(report["attendance_count"], 3)
        self.assertEqual(sum(row[3] for row in report["rows"]), 1)
        detailed = self.report("por-programa")
        self.assertEqual({row[0]: row[1:] for row in detailed["rows"]},
                         {"VOCA": [1, 2], "TANF-M": [1, 1], "Total consolidado sin duplicados": [1, 3]})
        self.assertEqual(sum(row[3] for row in self.report("participaciones")["rows"]), 3)
        self.assertNotIn("Solo expediente", str(self.report("bonafide")["rows"]))

    def test_program_and_date_filters_do_not_leak_other_participation(self):
        one = self.report(program_ids={self.tanf.program_id})
        self.assertEqual((one["unique_count"], one["attendance_count"]), (1, 1))
        day = self.report(start_date=date(2024, 2, 2), end_date=date(2024, 2, 2))
        self.assertEqual((day["unique_count"], day["attendance_count"]), (1, 1))
        empty = self.report(start_date=date(2024, 3, 1), end_date=date(2024, 3, 31))
        self.assertEqual((empty["unique_count"], empty["attendance_count"]), (0, 0))

    def test_discharge_keeps_attendance_in_historical_reports(self):
        before = self.report("bonafide")
        discharge_participant(self.db, participant_id=self.person.participant_id, program_id=self.tanf.program_id,
                              fiscal_year_id=self.year.fiscal_year_id, end_date=date(2024, 3, 1),
                              reason="Finalizó", actor_user_id=self.actor.user_id)
        after = self.report("bonafide")
        self.assertEqual((after["unique_count"], after["attendance_count"], after["rows"]),
                         (before["unique_count"], before["attendance_count"], before["rows"]))

    def test_frozen_snapshot_supplies_name_town_and_age_reference(self):
        set_snapshot_freeze(self.db, self.year.fiscal_year_id, frozen=True, actor_user_id=self.actor.user_id)
        self.person.nombre = "Nombre actual distinto"
        self.person.pueblo = "San Juan"
        self.person.fecha_nacimiento = date(1980, 1, 1)
        self.db.flush()
        row = self.report("bonafide")["rows"][0]
        self.assertIn("María", row[1])
        self.assertEqual(row[3:5], [14, "Ponce"])

    def test_adm_preserves_service_categories_per_program(self):
        report = self.report("adm")
        rows = report["rows"]
        self.assertEqual({row[0]: row[1:] for row in rows},
                         {"VOCA": ["Orientación", 2, 2, 1], "TANF-M": ["Orientación", 1, 1, 1]})
        demographics = report["sections"][0]
        self.assertEqual(demographics["headers"], ["Edad", "Femenino", "Masculino", "Total", "%"])
        self.assertTrue(all(len(row) == 5 for row in demographics["rows"]))
        self.assertEqual(sum(row[3] for row in demographics["rows"]), 1)

    def test_adm_unclassified_attendance_is_visible_but_excluded_from_service_totals(self):
        activity = create_activity(self.db, program_id=self.voca.program_id, code="SIN-ADM", description="Sin clasificación",
                                    fiscal_year_ids=[self.year.fiscal_year_id])
        session = create_activity_session(self.db, fiscal_year_id=self.year.fiscal_year_id,
            program_id=self.voca.program_id, activity_id=activity.activity_id,
            session_date=date(2024, 2, 4), actor_user_id=self.actor.user_id)
        set_session_attendance(self.db, session_id=session.session_id,
                                present_participant_ids=[self.person.participant_id, self.nonattender.participant_id])
        report = self.report("adm")
        self.assertEqual((report["unique_count"], report["attendance_count"]), (1, 3))
        unclassified = [section for section in report["sections"] if "sin clasificación" in section["title"].lower()]
        self.assertEqual(len(unclassified), 1)
        self.assertIn("SIN-ADM", str(unclassified[0]["rows"]))

    def test_school_results_retain_program_month_and_average(self):
        report = self.report("notas")
        self.assertEqual(len(report["rows"]), 1)
        row = report["rows"][0]
        self.assertEqual(row[:3], ["VOCA", "2024-02", "CP-2024-0001"])
        self.assertEqual(row[4:7], ["8", 90, 80])
        self.assertEqual(row[14:], [85, "B", "No"])

    def test_school_metrics_count_only_attendance_in_programs_with_academic_results(self):
        session = self.attend(self.tanf, date(2024, 2, 4))
        set_session_attendance(self.db, session_id=session.session_id,
                                present_participant_ids=[self.person.participant_id, self.nonattender.participant_id])
        report = self.report("notas")
        self.assertEqual((report["unique_count"], report["attendance_count"]), (1, 2))
        self.assertEqual(len(report["rows"]), 1)

    def test_missing_snapshot_is_explicit_error(self):
        missing = self.participant("Sin copia histórica")
        self.db.add(CPAttendance(session_id=self.session.session_id, participant_id=missing.participant_id, is_present=True))
        self.db.flush()
        with self.assertRaisesRegex(ValueError, "sin datos del año fiscal"):
            self.report()

    def test_six_report_types_generate_readable_pdf_and_excel(self):
        self.assertEqual(set(REPORT_TYPES), {"bonafide", "no-duplicados", "participaciones", "por-programa", "adm", "notas"})
        for report_type in REPORT_TYPES:
            with self.subTest(report_type=report_type):
                report = self.report(report_type)
                pdf = PdfReader(BytesIO(pdf_bytes(report)))
                self.assertGreaterEqual(len(pdf.pages), 1)
                text = " ".join(page.extract_text() for page in pdf.pages)
                self.assertIn("Comunidad y Prevención", text)
                self.assertNotIn("VCA", text)
                self.assertIn("Participantes no duplicados: 1", text)
                expected_attendances = 2 if report_type == "notas" else 3
                self.assertIn(f"Participaciones: {expected_attendances}", text)
                workbook = load_workbook(BytesIO(excel_bytes(report)), data_only=False)
                self.assertEqual(workbook.active["B3"].value, 1)
                self.assertEqual(workbook.active["D3"].value, expected_attendances)
                self.assertEqual(workbook.active["A5"].value, report["rows"][0][0])
                self.assertEqual(len(workbook.worksheets), 1 + len(report["sections"]))
                self.assertNotIn("VCA", [cell.value for sheet in workbook for row in sheet for cell in row])
                workbook.close()

    def test_excel_treats_formula_like_demographics_as_literal_text(self):
        self.person.pueblo = '=HYPERLINK("https://example.invalid","Texto")'
        sync_participants(self.db, self.year.fiscal_year_id, [self.person.participant_id], actor_user_id=self.actor.user_id)
        workbook = load_workbook(BytesIO(excel_bytes(self.report("bonafide"))), data_only=False)
        cell = workbook.active["E5"]
        self.assertEqual(cell.value, self.person.pueblo)
        self.assertEqual(cell.data_type, "s")
        self.assertFalse(any(cell.data_type == "f" for sheet in workbook for row in sheet for cell in row))
        workbook.close()


class CommunityReportRouteTests(_ReportFixture):
    def setUp(self):
        super().setUp()
        self.context = CommunityContext(self.actor, "viewer", (self.voca, self.tanf))
        app = FastAPI()
        app.add_middleware(SessionMiddleware, secret_key="reports-route-tests-not-production")
        app.mount("/static", StaticFiles(directory="app/static"), name="static")
        app.include_router(community_reports.router)
        app.dependency_overrides[get_db] = lambda: self.db
        app.dependency_overrides[require_community_context] = lambda: self.context
        self.client = TestClient(app, follow_redirects=False)

    def tearDown(self):
        self.client.close()
        super().tearDown()

    def params(self, **overrides):
        params = dict(fiscal_year_id=self.year.fiscal_year_id, program_ids=[self.voca.program_id],
                      report_type="bonafide", period_type="fiscal", output="screen")
        params.update(overrides)
        return params

    def test_viewer_can_read_download_and_see_available_programs(self):
        page = self.client.get("/community/reports")
        self.assertEqual(page.status_code, 200, page.text)
        self.assertIn("VOCA", page.text)
        self.assertIn("TANF-M", page.text)
        for report_type in REPORT_TYPES:
            for output, content_type in (("pdf", "application/pdf"), ("excel", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")):
                with self.subTest(report_type=report_type, output=output):
                    response = self.client.get("/community/reports", params=self.params(report_type=report_type, output=output))
                    self.assertEqual(response.status_code, 200, response.text[:100] if response.status_code != 200 else "")
                    self.assertEqual(response.headers["content-type"], content_type)
                    self.assertIn("attachment;", response.headers["content-disposition"])
                    if output == "pdf":
                        self.assertGreater(len(PdfReader(BytesIO(response.content)).pages), 0)
                    else:
                        workbook = load_workbook(BytesIO(response.content), read_only=True)
                        self.assertEqual(workbook.active["B3"].value, 1)
                        workbook.close()

    def test_user_program_scope_and_selected_context_enforced_for_all_outputs(self):
        self.context = CommunityContext(self.actor, "user", (self.voca, self.tanf), self.voca.program_id)
        for output in ("screen", "excel", "pdf"):
            with self.subTest(output=output):
                response = self.client.get("/community/reports", params=self.params(
                    program_ids=[self.voca.program_id, self.tanf.program_id], output=output))
                self.assertEqual(response.status_code, 403)
        page = self.client.get("/community/reports")
        self.assertNotIn("TANF-M", page.text)
        self.assertEqual(self.client.get("/community/reports", params=self.params()).status_code, 200)

    def test_dates_and_missing_selection_fail_without_download(self):
        invalid = self.client.get("/community/reports", params=self.params(output="excel", period_type="custom",
                                                                         start_date="2024-01-01", end_date="2024-02-28"))
        self.assertEqual(invalid.status_code, 422)
        self.assertEqual(self.client.get("/community/reports", params={"output": "pdf"}).status_code, 422)
        response = self.client.get("/community/reports", params=self.params(output="pdf", period_type="monthly", month=1, year=2024))
        self.assertEqual(response.status_code, 200)
        text = " ".join(page.extract_text() for page in PdfReader(BytesIO(response.content)).pages)
        self.assertIn("2024-01-15", text)
        self.assertIn("Participantes no duplicados: 0", text)


if __name__ == "__main__":
    unittest.main()
