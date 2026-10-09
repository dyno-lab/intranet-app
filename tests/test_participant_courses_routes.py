from copy import deepcopy
from io import BytesIO
import re
import unittest
from unittest.mock import patch

from fastapi import Depends, FastAPI, HTTPException, Request
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool
from starlette.middleware.sessions import SessionMiddleware

from tests import test_participant_courses as fixtures
from app.api.deps import get_db
from app.core.auth import get_current_user
from app.core.residential_scope import require_faro_access
from app.api.routes import participant_courses as routes
from app.api.routes import reports
from app.models.participant_monthly_course import ParticipantMonthlyCourse
from app.services.participant_courses_exports import course_report_pdf, course_report_excel


class ParticipantCourseRoutesTests(unittest.TestCase):
    def setUp(self):
        self.source = fixtures.ParticipantCoursesTests()
        with patch.object(fixtures.fixtures, 'create_engine', side_effect=lambda *_args: create_engine(
                'sqlite://', connect_args={'check_same_thread': False}, poolclass=StaticPool)):
            self.source.setUp()
        self.addCleanup(self.source.doCleanups)
        self.user = self.source.user
        self.app = FastAPI()
        self.app.add_middleware(SessionMiddleware, secret_key='courses-test-session-secret')
        self.app.include_router(routes.router, prefix='/ui/reports')
        self.app.include_router(reports.router, prefix='/ui/reports')
        def db():
            with Session(self.source.source.engine) as session:
                yield session
        def user():
            if self.user is None:
                raise HTTPException(303, headers={'Location': '/home'})
            return self.user
        self.app.dependency_overrides[get_db] = db
        self.app.dependency_overrides[get_current_user] = user
        self.client = TestClient(self.app)
        self.addCleanup(self.client.close)
        self.params = dict(proposal_id=[1, 2], month=7, year=2026, employee_id=0)

    def prepare(self):
        response = self.client.get('/ui/reports/cursos', params=self.params)
        self.assertEqual(response.status_code, 200, response.text)
        return re.search(r'data-token="([^"]+)"', response.text).group(1)

    def save(self, changes, token=None):
        return self.client.post('/ui/reports/cursos/save', params=self.params, json=changes,
                                headers={'X-CSRF-Token': token if token is not None else self.prepare()})

    def test_screen_save_and_reload_with_csrf_and_versions(self):
        response = self.save([self.source.change()])
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()['saved'][0]['revision'], 1)
        page = self.client.get('/ui/reports/cursos', params=self.params)
        self.assertIn('data-saved="reposteria"', page.text)
        self.assertIn('data-revision="1"', page.text)
        self.assertEqual(page.headers['cache-control'], 'no-store')
        self.assertEqual(self.save([self.source.change()], token='wrong').status_code, 403)
        self.assertEqual(self.save([self.source.change(course='charcuteria')]).status_code, 409)

    def test_auth_readonly_roles_and_scope_are_enforced_server_side(self):
        token = self.prepare()
        self.user.role = 'viewer'
        page = self.client.get('/ui/reports/cursos', params=self.params)
        self.assertNotIn('id="course-save"', page.text)
        self.assertEqual(self.save([self.source.change()], token=token).status_code, 403)
        self.user.role, self.user.residential_id = 'user', 2
        self.assertEqual(self.save([self.source.change(participant=2)], token=token).status_code, 403)
        self.user = None
        response = self.client.get('/ui/reports/cursos', params=self.params, follow_redirects=False)
        self.assertEqual(response.status_code, 303)

    def test_supervisor_sees_selects_and_can_save_with_csrf(self):
        self.user.role = 'supervisor'
        page = self.client.get('/ui/reports/cursos', params=self.params)
        self.assertIn('class="form-select course-choice"', page.text)
        self.assertIn('id="course-save"', page.text)
        self.assertEqual(self.save([self.source.change()], token='wrong').status_code, 403)
        saved = self.save([self.source.change()])
        self.assertEqual(saved.status_code, 200, saved.text)
        page = self.client.get('/ui/reports/cursos', params=self.params)
        self.assertIn('data-saved="reposteria"', page.text)
        record = self.source.db.scalars(select(ParticipantMonthlyCourse)).one()
        self.assertEqual(record.updated_by_user_id, self.user.user_id)

    def test_assigned_residential_context_allows_courses_and_blocks_other_residentials(self):
        # Exercise the same assignment dependency as main.py, including users
        # whose legacy residential_id differs from their selected assignment.
        self.user.role, self.user.residential_id = 'user', 1
        for rid in (1, 2):
            self.source.source.insert('user_residentials', user_residential_id=rid,
                                      user_id=self.user.user_id, residential_id=rid, is_active=True)
        self.source.db.commit()
        app = FastAPI()
        app.add_middleware(SessionMiddleware, secret_key='courses-test-session-secret')
        app.include_router(routes.router, prefix='/ui/reports', dependencies=[Depends(require_faro_access)])
        app.dependency_overrides.update(self.app.dependency_overrides)

        @app.get('/test/residential/{rid}')
        def select_residential(request: Request, rid: int):
            request.session['active_residential_id'] = rid

        with patch('app.core.residential_scope._FARO_PERMISSION_DEPENDENCY', return_value=self.user), TestClient(app) as client:
            for role in ('user', 'supervisor'):
                self.user.role = role
                with self.subTest(role=role):
                    client.get('/test/residential/2')
                    # A crafted report filter cannot override the assignment.
                    params = {**self.params, 'employee_id': -1}
                    page = client.get('/ui/reports/cursos', params=params)
                    self.assertEqual(page.status_code, 200, page.text)
                    self.assertIn('data-participant="1"', page.text)
                    self.assertNotIn('data-participant="2"', page.text)
                    token = re.search(r'data-token="([^"]+)"', page.text).group(1)
                    revision = int(re.search(r'data-revision="(\d+)"', page.text).group(1))
                    saved = client.post('/ui/reports/cursos/save', params=params,
                        json=[self.source.change(revision=revision)], headers={'X-CSRF-Token': token})
                    self.assertEqual(saved.status_code, 200, saved.text)
                    outside = client.post('/ui/reports/cursos/save', params=params,
                        json=[self.source.change(participant=2)], headers={'X-CSRF-Token': token})
                    self.assertEqual(outside.status_code, 403)
                    client.get('/test/residential/1')
                    page = client.get('/ui/reports/cursos', params=self.params)
                    self.assertIn('data-participant="2"', page.text)
                    client.get('/test/residential/999')
                    denied = client.get('/ui/reports/cursos', params=self.params, follow_redirects=False)
                    self.assertEqual(denied.status_code, 303)
                    self.assertTrue(denied.headers['location'].startswith('/login?'))

    def test_all_exports_use_saved_choices_and_same_people(self):
        self.save([self.source.change()])
        for suffix, disposition in [('/pdf', 'inline'), ('/pdf/download', 'attachment')]:
            response = self.client.get('/ui/reports/cursos' + suffix, params=self.params)
            self.assertEqual(response.status_code, 200, response.text[:200] if response.status_code != 200 else '')
            self.assertTrue(response.headers['content-disposition'].startswith(disposition))
            reader = PdfReader(BytesIO(response.content))
            body = '\n'.join(p.extract_text() for p in reader.pages)
            self.assertIn('Repostería', body)
            self.assertIn('Participantes únicos en todo el período: 2', body)
            self.assertEqual(len(reader.pages[0].images), 2)
        response = self.client.get('/ui/reports/cursos/excel', params=self.params)
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet['D9'].value, 'Repostería')
        self.assertEqual(sheet['B12'].value, 2)
        self.assertEqual(sheet.freeze_panes, 'D9')

    def test_catalog_redirect_preserves_all_proposals_custom_dates_and_global_zero(self):
        for output, suffix in [('screen', ''), ('pdf', '/pdf'), ('excel', '/excel')]:
            response = self.client.get('/ui/reports/run', params={**self.params, 'report_key': 'cursos',
                'period_type': 'custom', 'start_date': '2026-07-01', 'end_date': '2026-08-31', 'output': output}, follow_redirects=False)
            self.assertEqual(response.status_code, 303)
            url = response.headers['location']
            self.assertIn('/cursos' + suffix + '?', url)
            for text in ['proposal_id=1', 'proposal_id=2', 'employee_id=0', 'end_date=2026-08-31']:
                self.assertIn(text, url)

    def test_malformed_and_oversized_save_preserve_database(self):
        token = self.prepare()
        for content, code in [('not-json', 422), (' ' * (1024 * 1024 + 1), 413)]:
            response = self.client.post('/ui/reports/cursos/save', params=self.params, content=content, headers={'X-CSRF-Token': token})
            self.assertEqual(response.status_code, code)
        self.assertEqual(self.source.db.scalars(select(ParticipantMonthlyCourse)).all(), [])

    def test_pdf_continuations_empty_report_and_excel_formula_safety(self):
        data = self.source.build(period_type='custom', start_date='2026-06-01', end_date='2026-09-30')
        original = data['rows'][0]
        data['rows'] = []
        for i in range(70):
            row = deepcopy(original)
            row['name'] = f'Persona de prueba {i:03}'
            data['rows'].append(row)
        data['total'] = 70
        reader = PdfReader(BytesIO(course_report_pdf(data)))
        self.assertGreater(len(reader.pages), 2)
        body = '\n'.join(page.extract_text() for page in reader.pages)
        self.assertEqual(body.count('Persona de prueba 069'), 2)  # two month groups, one global count
        self.assertIn('Septiembre 2026', body)
        for page in reader.pages:
            self.assertEqual(len(page.images), 2)
            self.assertIn('Informe de cursos', page.extract_text())
        data['rows'][0]['name'] = '=1+1'
        sheet = load_workbook(BytesIO(course_report_excel(data))).active
        self.assertEqual(sheet['A9'].data_type, 's')
        self.assertEqual(sheet['A9'].value, '=1+1')
        empty = self.source.build(month=9)
        self.assertIn('Sin participantes', PdfReader(BytesIO(course_report_pdf(empty))).pages[0].extract_text())

    def test_pdf_total_never_creates_a_page_without_participants(self):
        data = self.source.build(period_type='custom', start_date='2026-06-01', end_date='2026-09-30')
        original = data['rows'][0]
        data['rows'] = [deepcopy(original) for _ in range(9)]
        for index, row in enumerate(data['rows']):
            row['name'] = f'Persona de prueba {index:03}'
        data['total'] = 9
        for page in PdfReader(BytesIO(course_report_pdf(data))).pages:
            self.assertIn('Persona de prueba', page.extract_text())
